"""
technical_analysis.py  -  Phân tích kỹ thuật cho NHIỀU MÃ cổ phiếu cùng lúc.

Đầu vào: file CSV/XLSX dạng "dài" (long) do data_loader.build_stock_prices_csv tạo ra:
    Date, Symbol, Open, High, Low, Close, Volume      (mỗi dòng = 1 mã x 1 phiên)

Cài thư viện:
    pip install pandas numpy openpyxl

Cách chạy:
    python technical_analysis.py stock_prices.csv
    python technical_analysis.py stock_prices.csv --symbols FPT VCB VRE
    python technical_analysis.py stock_prices.csv --verify          # đối chiếu công thức
    python technical_analysis.py stock_prices.csv --keep-last 0     # lưu toàn bộ chỉ báo (file rất lớn)

Kết quả trong thư mục output_technical/:
    technical_report.xlsx          Excel gồm: Tổng hợp, Tín hiệu chi tiết, Kiểm tra dữ liệu, Thống kê
    technical_summary.csv          mỗi mã một dòng: chỉ báo + tín hiệu + nhận định
    technical_signals_long.csv     mỗi mã x mỗi chỉ báo một dòng, có giải thích
    data_quality_by_symbol.csv     kiểm tra dữ liệu theo từng mã
    technical_indicators_lastN.csv chỉ báo theo phiên (N phiên gần nhất mỗi mã)
    latest_assessment.txt          nhận định dạng văn bản cho từng mã

Quy ước công thức (chuẩn phổ biến, khớp TradingView):
    SMA20/SMA50 = trung bình đơn giản 20/50 phiên
    EMA12/EMA26 = trung bình động hàm mũ, alpha = 2/(span+1), adjust=False
    RSI14       = Wilder, alpha = 1/14 (KHÁC EMA thường)
    MACD        = EMA12 - EMA26; Signal = EMA9 của MACD; Histogram = MACD - Signal
    Bollinger   = SMA20 +/- 2 độ lệch chuẩn tổng thể (ddof=0) của 20 phiên
Mọi chỉ báo được tính RIÊNG cho từng mã, không bao giờ trộn giá giữa hai mã.
"""

from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd

VN_TZ = "Asia/Ho_Chi_Minh"
MIN_ROWS = 60        # dưới ngưỡng này: không kết luận tín hiệu (cần >= 50 phiên cho SMA50)
STABLE_ROWS = 150    # dưới ngưỡng này: EMA/MACD/RSI chưa ổn định hoàn toàn
BIG_MOVE = 0.15      # biến động 1 phiên > 15%: nghi giá chưa điều chỉnh / lỗi dữ liệu
STALE_SESSIONS = 5   # mã dừng cập nhật quá số phiên này so với thị trường thì cảnh báo

PRICE_COLS = ["Open", "High", "Low", "Close"]

COLUMN_ALIASES = {
    "Date": ["date", "datetime", "tradingdate", "trading date", "time", "ngày", "ngay", "thời gian", "thoi gian"],
    "Symbol": ["symbol", "ticker", "code", "mã", "ma", "mã ck", "ma ck", "mã cp", "ma cp"],
    "Open": ["open", "o", "giá mở cửa", "gia mo cua"],
    "High": ["high", "h", "giá cao nhất", "gia cao nhat"],
    "Low": ["low", "l", "giá thấp nhất", "gia thap nhat"],
    "Close": ["close", "c", "giá đóng cửa", "gia dong cua", "closing price", "price", "adj close"],
    "Volume": ["volume", "vol", "v", "khối lượng", "khoi luong", "kl", "totalvolume"],
}


# =====================================================================
# 1. ĐỌC DỮ LIỆU
# =====================================================================
def _norm(name: object) -> str:
    return str(name).strip().lower().replace("-", " ").replace("_", " ")


def _parse_dates(s: pd.Series) -> pd.Series:
    """ISO (2026-10-09) là chính; nếu không đọc được thì thử dd/mm/yyyy; số nguyên lớn = epoch."""
    if pd.api.types.is_datetime64_any_dtype(s):
        return s.dt.tz_localize(None).dt.normalize() if getattr(s.dt, "tz", None) else s.dt.normalize()
    if pd.api.types.is_numeric_dtype(s) and s.notna().any():
        med = float(s.dropna().median())
        unit = "ms" if med > 1e11 else ("s" if med > 1e8 else None)
        if unit:
            return (pd.to_datetime(s, unit=unit, utc=True).dt.tz_convert(VN_TZ)
                    .dt.tz_localize(None).dt.normalize())
    txt = s.astype(str).str.strip()
    out = pd.to_datetime(txt, format="ISO8601", errors="coerce")
    bad = out.isna() & ~txt.str.lower().isin(["", "nan", "none", "nat"])
    if bad.any():
        out.loc[bad] = pd.to_datetime(txt[bad], dayfirst=True, format="mixed", errors="coerce")
    return out.dt.normalize()


def _to_num(s: pd.Series) -> pd.Series:
    if pd.api.types.is_numeric_dtype(s):
        return pd.to_numeric(s, errors="coerce")
    return pd.to_numeric(s.astype(str).str.replace(",", "", regex=False).str.strip(), errors="coerce")


def load_data(file_path: str | Path, default_symbol: str | None = None) -> pd.DataFrame:
    """Đọc CSV/XLSX nhiều mã -> DataFrame cột Date, Symbol, Open, High, Low, Close, Volume.

    Sắp xếp theo (Symbol, Date); trùng (Symbol, Date) giữ dòng cuối.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Không tìm thấy file dữ liệu: {path}")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        try:
            raw = pd.read_csv(path, encoding="utf-8-sig")
        except UnicodeDecodeError:
            raw = pd.read_csv(path, encoding="cp1258")
    elif suffix in (".xlsx", ".xls"):
        raw = pd.read_excel(path)
    else:
        raise ValueError("Chỉ hỗ trợ file .csv, .xlsx hoặc .xls")

    raw.columns = [str(c).strip() for c in raw.columns]
    present = {_norm(c): c for c in raw.columns}
    rename = {}
    for std, aliases in COLUMN_ALIASES.items():
        found = next((present[_norm(a)] for a in aliases if _norm(a) in present), None)  # ưu tiên theo thứ tự alias
        if found is not None:
            rename[found] = std
    df = raw.rename(columns=rename)

    if "Symbol" not in df.columns:
        sym = (default_symbol or path.stem).upper()
        print(f"[CẢNH BÁO] File không có cột mã (Symbol): coi toàn bộ là mã '{sym}'.")
        df["Symbol"] = sym

    missing = [c for c in ["Date", "Open", "High", "Low", "Close", "Volume"] if c not in df.columns]
    if missing:
        raise ValueError(f"Thiếu cột bắt buộc: {missing}. Các cột hiện có: {list(raw.columns)}")

    df = df[["Date", "Symbol", "Open", "High", "Low", "Close", "Volume"]].copy()
    df["Date"] = _parse_dates(df["Date"])
    for c in PRICE_COLS + ["Volume"]:
        df[c] = _to_num(df[c])
    # Làm tròn bỏ nhiễu dấu phẩy động (ví dụ 32159.999999999996 -> 32160.0).
    df[PRICE_COLS] = df[PRICE_COLS].round(2)
    df["Symbol"] = df["Symbol"].astype(str).str.strip().str.upper()

    df = df.dropna(subset=["Date"])
    df = df[(df["Symbol"] != "") & (df["Symbol"] != "NAN")]
    n_before = len(df)
    df = (df.sort_values(["Symbol", "Date"])
            .drop_duplicates(subset=["Symbol", "Date"], keep="last")
            .reset_index(drop=True))
    df.attrs["duplicates_removed"] = n_before - len(df)
    return df


# =====================================================================
# 2. KIỂM TRA DỮ LIỆU THEO TỪNG MÃ
# =====================================================================
def check_data(df: pd.DataFrame) -> pd.DataFrame:
    """Mỗi mã một dòng. Lịch giao dịch chung = các ngày có >= 50% số mã giao dịch."""
    n_sym = df["Symbol"].nunique()
    per_day = df.groupby("Date")["Symbol"].nunique()
    calendar = per_day.index[per_day >= max(1, 0.5 * n_sym)]

    rows = []
    for sym, d in df.groupby("Symbol", sort=True):
        d = d.sort_values("Date")
        first, last = d["Date"].iloc[0], d["Date"].iloc[-1]
        expected = calendar[(calendar >= first) & (calendar <= last)]
        missing_sessions = int(len(expected.difference(pd.DatetimeIndex(d["Date"]))))
        behind = int((calendar > last).sum())

        nan_cells = int(d[PRICE_COLS + ["Volume"]].isna().sum().sum())
        nonpos = int((d[PRICE_COLS] <= 0).any(axis=1).sum())
        neg_vol = int((d["Volume"] < 0).sum())
        zero_vol = int((d["Volume"] == 0).sum())
        hi = d[["Open", "Low", "Close"]].max(axis=1)
        lo = d[["Open", "High", "Close"]].min(axis=1)
        bad_ohlc = int(((d["High"] < hi) | (d["Low"] > lo)).sum())
        ret = d["Close"].where(d["Close"] > 0).pct_change().abs()
        max_move = float(ret.max()) if ret.notna().any() else 0.0
        big_moves = int((ret > BIG_MOVE).sum())
        n = len(d)

        warns = []
        if n < MIN_ROWS:
            warns.append(f"Chỉ có {n} phiên (<{MIN_ROWS}): không đủ tính SMA50/RSI/MACD")
        elif n < STABLE_ROWS:
            warns.append(f"Chỉ có {n} phiên (<{STABLE_ROWS}): EMA/MACD/RSI chưa ổn định hoàn toàn")
        if behind >= STALE_SESSIONS:
            warns.append(f"Dữ liệu dừng từ {last:%d/%m/%Y}, thiếu {behind} phiên so với thị trường "
                         "(tạm ngừng/hủy niêm yết/chưa cập nhật)")
        if missing_sessions:
            warns.append(f"Thiếu {missing_sessions} phiên so với lịch giao dịch chung")
        if nan_cells:
            warns.append(f"{nan_cells} ô thiếu giá trị")
        if nonpos:
            warns.append(f"{nonpos} phiên có giá <= 0")
        if neg_vol:
            warns.append(f"{neg_vol} phiên khối lượng âm")
        if bad_ohlc:
            warns.append(f"{bad_ohlc} phiên OHLC sai logic")
        if big_moves:
            warns.append(f"{big_moves} phiên biến động >{BIG_MOVE:.0%}: có thể giá chưa điều chỉnh khi chia tách/cổ tức")
        if n and zero_vol / n > 0.2:
            warns.append(f"{zero_vol} phiên khối lượng = 0 (thanh khoản rất thấp/ngừng giao dịch)")

        rows.append({
            "Mã": sym, "Số phiên": n, "Ngày đầu": first, "Ngày cuối": last,
            "Phiên thiếu so với lịch chung": missing_sessions,
            "Số phiên thị trường sau ngày cuối": behind,
            "Ô thiếu (NaN)": nan_cells, "Giá <= 0": nonpos, "Khối lượng âm": neg_vol,
            "Khối lượng = 0": zero_vol, "OHLC sai logic": bad_ohlc,
            "Biến động 1 phiên lớn nhất (%)": round(max_move * 100, 2),
            f"Số phiên biến động >{BIG_MOVE:.0%}": big_moves,
            "Chất lượng dữ liệu": "Không đủ" if n < MIN_ROWS else ("Cảnh báo" if warns else "Tốt"),
            "Cảnh báo dữ liệu": "; ".join(warns),
        })
    return pd.DataFrame(rows)


# =====================================================================
# 3. CHỈ BÁO (tính riêng từng mã)
# =====================================================================
def _wilder_rsi(close: pd.Series, n: int = 14) -> pd.Series:
    delta = close.diff()
    gain = delta.clip(lower=0)
    loss = -delta.clip(upper=0)
    avg_gain = gain.ewm(alpha=1 / n, adjust=False, min_periods=n).mean()
    avg_loss = loss.ewm(alpha=1 / n, adjust=False, min_periods=n).mean()
    rs = avg_gain / avg_loss.replace(0, np.nan)
    rsi = 100 - 100 / (1 + rs)
    rsi = rsi.where(~((avg_loss == 0) & (avg_gain > 0)), 100.0)   # chỉ tăng -> 100
    rsi = rsi.where(~((avg_loss == 0) & (avg_gain == 0)), 50.0)   # đứng giá -> 50
    return rsi


def add_indicators_one(d: pd.DataFrame) -> pd.DataFrame:
    """d: dữ liệu MỘT mã, đã sắp xếp theo ngày."""
    out = d.copy()
    close = out["Close"]
    out["SMA20"] = close.rolling(20, min_periods=20).mean()
    out["SMA50"] = close.rolling(50, min_periods=50).mean()
    out["EMA12"] = close.ewm(span=12, adjust=False, min_periods=12).mean()
    out["EMA26"] = close.ewm(span=26, adjust=False, min_periods=26).mean()
    out["RSI14"] = _wilder_rsi(close, 14)
    out["MACD"] = out["EMA12"] - out["EMA26"]
    out["MACD_Signal"] = out["MACD"].ewm(span=9, adjust=False, min_periods=9).mean()
    out["MACD_Hist"] = out["MACD"] - out["MACD_Signal"]
    std20 = close.rolling(20, min_periods=20).std(ddof=0)
    out["BB_Middle"] = out["SMA20"]
    out["BB_Upper"] = out["BB_Middle"] + 2 * std20
    out["BB_Lower"] = out["BB_Middle"] - 2 * std20
    return out


def add_indicators(df: pd.DataFrame) -> pd.DataFrame:
    """Tính chỉ báo cho từng mã rồi ghép lại. Bỏ dòng thiếu/<=0 giá đóng cửa."""
    parts = []
    for _, d in df.groupby("Symbol", sort=True):
        d = d.dropna(subset=["Close"])
        d = d[d["Close"] > 0].sort_values("Date").reset_index(drop=True)
        if not d.empty:
            parts.append(add_indicators_one(d))
    if not parts:
        raise ValueError("Không có mã nào có giá đóng cửa hợp lệ.")
    return pd.concat(parts, ignore_index=True)


# =====================================================================
# 4. TÍN HIỆU (quy tắc giữ nguyên bản gốc)
# =====================================================================
def classify_symbol(d: pd.DataFrame, min_rows: int = MIN_ROWS) -> tuple[list[dict], dict]:
    """d: chỉ báo của MỘT mã. Trả về (danh sách tín hiệu từng chỉ báo, dòng tổng hợp)."""
    r = d.iloc[-1]
    prev = d.iloc[-2] if len(d) > 1 else None
    n = len(d)
    sym = str(r["Symbol"])
    signals: list[dict] = []

    def add(name, value, signal, reason):
        signals.append({"Mã": sym, "Chỉ báo": name, "Giá trị phiên gần nhất": value,
                        "Tín hiệu": signal, "Giải thích": reason})

    # --- Xu hướng: SMA20/SMA50 ---
    if pd.notna(r["SMA20"]) and pd.notna(r["SMA50"]):
        v = f"SMA20={r['SMA20']:,.0f}; SMA50={r['SMA50']:,.0f}"
        if r["Close"] > r["SMA20"] > r["SMA50"]:
            add("SMA20/SMA50", v, "Tích cực", "Giá trên SMA20 và SMA20 trên SMA50, xu hướng tăng được xác nhận tương đối.")
        elif r["Close"] < r["SMA20"] < r["SMA50"]:
            add("SMA20/SMA50", v, "Tiêu cực", "Giá dưới SMA20 và SMA20 dưới SMA50, xu hướng giảm chiếm ưu thế.")
        else:
            add("SMA20/SMA50", v, "Trung lập", "Vị trí giá và hai đường SMA chưa đồng thuận rõ ràng.")
    else:
        add("SMA20/SMA50", "Chưa đủ dữ liệu", "Trung lập", "Cần ít nhất 50 phiên để tính đủ SMA50.")

    # --- RSI ---
    if pd.notna(r["RSI14"]):
        x = float(r["RSI14"])
        if x >= 70:
            s, why = "Tiêu cực", "RSI từ 70 trở lên: vùng quá mua, có thể xuất hiện điều chỉnh; không đồng nghĩa chắc chắn sẽ giảm."
        elif x <= 30:
            s, why = "Tích cực", "RSI từ 30 trở xuống: vùng quá bán, có thể xuất hiện hồi phục; cần xác nhận thêm."
        elif x >= 55:
            s, why = "Tích cực", "RSI trên 55 cho thấy động lượng tăng tương đối mạnh."
        elif x <= 45:
            s, why = "Tiêu cực", "RSI dưới 45 cho thấy động lượng yếu tương đối."
        else:
            s, why = "Trung lập", "RSI nằm trong vùng giữa, động lượng chưa nghiêng rõ về bên mua hoặc bên bán."
        add("RSI14", round(x, 2), s, why)
    else:
        add("RSI14", "Chưa đủ dữ liệu", "Trung lập", "Cần thêm dữ liệu lịch sử để tính RSI14.")

    # --- MACD ---
    if pd.notna(r["MACD"]) and pd.notna(r["MACD_Signal"]):
        v = f"MACD={r['MACD']:,.2f}; Signal={r['MACD_Signal']:,.2f}; Hist={r['MACD_Hist']:,.2f}"
        if r["MACD"] > r["MACD_Signal"] and r["MACD_Hist"] > 0:
            s, why = "Tích cực", "MACD nằm trên đường tín hiệu và histogram dương, động lượng tăng đang chiếm ưu thế."
        elif r["MACD"] < r["MACD_Signal"] and r["MACD_Hist"] < 0:
            s, why = "Tiêu cực", "MACD nằm dưới đường tín hiệu và histogram âm, động lượng giảm đang chiếm ưu thế."
        else:
            s, why = "Trung lập", "MACD và đường tín hiệu chưa cho tín hiệu đồng thuận mạnh."
        add("MACD", v, s, why)
    else:
        add("MACD", "Chưa đủ dữ liệu", "Trung lập", "Cần thêm dữ liệu lịch sử để tính MACD và đường tín hiệu.")

    # --- Bollinger Bands ---
    if pd.notna(r["BB_Upper"]) and pd.notna(r["BB_Lower"]):
        v = f"Upper={r['BB_Upper']:,.0f}; Middle={r['BB_Middle']:,.0f}; Lower={r['BB_Lower']:,.0f}"
        if r["Close"] > r["BB_Upper"]:
            s, why = "Trung lập", "Giá vượt dải trên: có thể là đà tăng mạnh hoặc trạng thái quá xa trung bình; cần kết hợp chỉ báo khác."
        elif r["Close"] < r["BB_Lower"]:
            s, why = "Trung lập", "Giá dưới dải dưới: có thể là áp lực bán mạnh hoặc trạng thái quá xa trung bình; cần xác nhận thêm."
        elif r["Close"] >= r["BB_Middle"]:
            s, why = "Tích cực", "Giá ở nửa trên dải Bollinger, thiên hướng ngắn hạn tương đối tích cực."
        else:
            s, why = "Tiêu cực", "Giá ở nửa dưới dải Bollinger, thiên hướng ngắn hạn tương đối yếu."
        add("Bollinger Bands", v, s, why)
    else:
        add("Bollinger Bands", "Chưa đủ dữ liệu", "Trung lập", "Cần ít nhất 20 phiên để tính dải Bollinger.")

    sig = {s["Chỉ báo"]: s["Tín hiệu"] for s in signals}
    pos = sum(v == "Tích cực" for v in sig.values())
    neu = sum(v == "Trung lập" for v in sig.values())
    neg = sum(v == "Tiêu cực" for v in sig.values())

    if n < min_rows:
        overall = "CHƯA ĐỦ DỮ LIỆU"
        explanation = f"Chỉ có {n} phiên (<{min_rows}); không đủ để kết luận tín hiệu."
    elif pos >= neg + 2:
        overall = "TÍCH CỰC"
        explanation = "Số tín hiệu tích cực nhiều hơn đáng kể tín hiệu tiêu cực. Đây là nhận định kỹ thuật, không phải khuyến nghị mua bán."
    elif neg >= pos + 2:
        overall = "TIÊU CỰC"
        explanation = "Số tín hiệu tiêu cực nhiều hơn đáng kể tín hiệu tích cực. Cần quản trị rủi ro và xác nhận bằng dữ liệu khác."
    else:
        overall = "TRUNG LẬP / TÍN HIỆU CHƯA ĐỒNG THUẬN"
        explanation = "Các chỉ báo chưa tạo được chênh lệch đủ lớn để kết luận xu hướng rõ ràng."

    # --- Xu hướng & động lượng (nhãn tóm tắt) ---
    trend = {"Tích cực": "Tăng", "Tiêu cực": "Giảm", "Trung lập": "Đi ngang / chưa rõ"}[sig["SMA20/SMA50"]]
    if pd.isna(r["SMA50"]):
        trend = "Chưa đủ dữ liệu"
    if pd.notna(r["RSI14"]) and pd.notna(r["MACD_Hist"]):
        if r["RSI14"] > 50 and r["MACD_Hist"] > 0:
            momentum = "Tăng (RSI>50, histogram>0)"
        elif r["RSI14"] < 50 and r["MACD_Hist"] < 0:
            momentum = "Giảm (RSI<50, histogram<0)"
        else:
            momentum = "Phân hóa / trung tính"
    else:
        momentum = "Chưa đủ dữ liệu"
    rsi_zone = ("Chưa đủ dữ liệu" if pd.isna(r["RSI14"]) else
                "Quá mua (>=70)" if r["RSI14"] >= 70 else
                "Quá bán (<=30)" if r["RSI14"] <= 30 else "Bình thường")

    # --- Sự kiện ở phiên cuối ---
    events = []
    if prev is not None:
        if pd.notna(prev["MACD"]) and pd.notna(prev["MACD_Signal"]) and pd.notna(r["MACD_Signal"]):
            if prev["MACD"] <= prev["MACD_Signal"] and r["MACD"] > r["MACD_Signal"]:
                events.append("MACD cắt lên đường tín hiệu")
            elif prev["MACD"] >= prev["MACD_Signal"] and r["MACD"] < r["MACD_Signal"]:
                events.append("MACD cắt xuống đường tín hiệu")
        if pd.notna(prev["SMA50"]) and pd.notna(r["SMA50"]):
            if prev["SMA20"] <= prev["SMA50"] and r["SMA20"] > r["SMA50"]:
                events.append("SMA20 cắt lên SMA50")
            elif prev["SMA20"] >= prev["SMA50"] and r["SMA20"] < r["SMA50"]:
                events.append("SMA20 cắt xuống SMA50")
    if pd.notna(r["BB_Upper"]):
        if r["Close"] > r["BB_Upper"]:
            events.append("Giá vượt dải Bollinger trên")
        elif r["Close"] < r["BB_Lower"]:
            events.append("Giá thủng dải Bollinger dưới")

    summary = {
        "Mã": sym, "Ngày dữ liệu cuối": r["Date"], "Số phiên": n, "Giá đóng cửa": r["Close"],
        "SMA20": r["SMA20"], "SMA50": r["SMA50"], "EMA12": r["EMA12"], "EMA26": r["EMA26"],
        "RSI14": r["RSI14"], "MACD": r["MACD"], "MACD_Signal": r["MACD_Signal"], "MACD_Hist": r["MACD_Hist"],
        "BB_Upper": r["BB_Upper"], "BB_Middle": r["BB_Middle"], "BB_Lower": r["BB_Lower"],
        "Tín hiệu SMA": sig["SMA20/SMA50"], "Tín hiệu RSI": sig["RSI14"],
        "Tín hiệu MACD": sig["MACD"], "Tín hiệu Bollinger": sig["Bollinger Bands"],
        "Số tích cực": pos, "Số trung lập": neu, "Số tiêu cực": neg,
        "Nhận định tổng hợp": overall, "Xu hướng": trend, "Động lượng": momentum, "Vùng RSI": rsi_zone,
        "Sự kiện phiên cuối": "; ".join(events),
        "_explanation": explanation,
    }
    return signals, summary


def build_assessment_text(summary: dict) -> str:
    return (
        f"[{summary['Mã']}] Ngày gần nhất: {pd.Timestamp(summary['Ngày dữ liệu cuối']).date()}\n"
        f"  Giá đóng cửa: {summary['Giá đóng cửa']:,.0f} đ\n"
        f"  Nhận định tổng hợp: {summary['Nhận định tổng hợp']}\n"
        f"  Số tín hiệu: tích cực={summary['Số tích cực']}, trung lập={summary['Số trung lập']}, tiêu cực={summary['Số tiêu cực']}\n"
        f"  Xu hướng: {summary['Xu hướng']} | Động lượng: {summary['Động lượng']} | RSI: {summary['Vùng RSI']}\n"
        f"  Giải thích: {summary['_explanation']}\n"
    )


def analyze_all(ind: pd.DataFrame, quality: pd.DataFrame, min_rows: int = MIN_ROWS):
    """Chạy phân loại cho từng mã. Trả về (bảng tổng hợp, bảng tín hiệu dài, văn bản nhận định)."""
    q = quality.set_index("Mã")
    summaries, longs, texts = [], [], []
    for sym, d in ind.groupby("Symbol", sort=True):
        d = d.sort_values("Date").reset_index(drop=True)
        sigs, s = classify_symbol(d, min_rows)
        s["Chất lượng dữ liệu"] = q.loc[sym, "Chất lượng dữ liệu"] if sym in q.index else ""
        s["Cảnh báo dữ liệu"] = q.loc[sym, "Cảnh báo dữ liệu"] if sym in q.index else ""
        texts.append(build_assessment_text(s))
        s.pop("_explanation")
        summaries.append(s)
        longs.extend(sigs)
    return pd.DataFrame(summaries), pd.DataFrame(longs), "\n".join(texts)


def market_overview(summary: pd.DataFrame, quality: pd.DataFrame) -> pd.DataFrame:
    rows = [("Số mã phân tích", len(summary))]
    for label, col in [("Nhận định tổng hợp", "Nhận định tổng hợp"), ("Xu hướng", "Xu hướng"),
                       ("Động lượng", "Động lượng"), ("Vùng RSI", "Vùng RSI")]:
        for k, v in summary[col].value_counts().items():
            rows.append((f"{label}: {k}", int(v)))
    for k, v in quality["Chất lượng dữ liệu"].value_counts().items():
        rows.append((f"Chất lượng dữ liệu: {k}", int(v)))
    return pd.DataFrame(rows, columns=["Chỉ tiêu", "Số mã"])


# =====================================================================
# 5. ĐỐI CHIẾU CÔNG THỨC (cài đặt độc lập bằng vòng lặp thuần)
# =====================================================================
def verify_formulas(d: pd.DataFrame) -> dict:
    """So sánh chỉ báo vector hóa với bản tính tay bằng vòng lặp cho MỘT mã.
    Trả về {tên: (sai số tuyệt đối lớn nhất, đạt/không)}."""
    c = d["Close"].to_numpy(dtype=float)
    n = len(c)
    res: dict[str, float] = {}

    ref = np.full(n, np.nan)
    for i in range(19, n):
        ref[i] = c[i - 19:i + 1].mean()
    res["SMA20"] = np.nanmax(np.abs(ref - d["SMA20"].to_numpy()))

    ref = np.full(n, np.nan)
    for i in range(49, n):
        ref[i] = c[i - 49:i + 1].mean()
    res["SMA50"] = np.nanmax(np.abs(ref - d["SMA50"].to_numpy())) if n >= 50 else np.nan

    def ema_loop(x, span):
        a = 2 / (span + 1)
        e = np.empty(len(x))
        e[0] = x[0]
        for i in range(1, len(x)):
            e[i] = a * x[i] + (1 - a) * e[i - 1]
        return e

    e12, e26 = ema_loop(c, 12), ema_loop(c, 26)
    macd = e12 - e26
    res["EMA12"] = np.nanmax(np.abs(e12 - d["EMA12"].to_numpy()))
    res["EMA26"] = np.nanmax(np.abs(e26 - d["EMA26"].to_numpy()))
    res["MACD"] = np.nanmax(np.abs(macd - d["MACD"].to_numpy()))
    # Signal = EMA9 của MACD, bắt đầu từ MACD hợp lệ đầu tiên (phiên thứ 26, chỉ số 25)
    # và chỉ có giá trị sau 9 quan sát (chỉ số 33), đúng như cách pandas ewm(min_periods=9) làm.
    sig = np.full(n, np.nan)
    if n > 34:
        sig[25] = macd[25]
        for i in range(26, n):
            sig[i] = 0.2 * macd[i] + 0.8 * sig[i - 1]
        res["MACD_Signal"] = np.nanmax(np.abs(sig[33:] - d["MACD_Signal"].to_numpy()[33:]))
    else:
        res["MACD_Signal"] = np.nan

    ref = np.full(n, np.nan)
    for i in range(19, n):
        w = c[i - 19:i + 1]
        ref[i] = w.mean() + 2 * np.sqrt(((w - w.mean()) ** 2).mean())
    res["BB_Upper"] = np.nanmax(np.abs(ref - d["BB_Upper"].to_numpy()))

    # RSI Wilder chuẩn: khởi tạo bằng trung bình đơn giản 14 phiên đầu rồi làm mượt.
    # Bản pandas khởi tạo khác ở vài chục phiên đầu nên chỉ so sánh sau 200 phiên.
    delta = np.diff(c)
    g, l = np.clip(delta, 0, None), np.clip(-delta, 0, None)
    rsi_ref = np.full(n, np.nan)
    if n > 15:
        ag, al = g[:14].mean(), l[:14].mean()
        rsi_ref[14] = 100 if al == 0 else 100 - 100 / (1 + ag / al)
        for i in range(14, len(delta)):
            ag = (ag * 13 + g[i]) / 14
            al = (al * 13 + l[i]) / 14
            rsi_ref[i + 1] = 100 if al == 0 else 100 - 100 / (1 + ag / al)
    res["RSI14 (sau 200 phiên)"] = (np.nanmax(np.abs(rsi_ref - d["RSI14"].to_numpy())[200:])
                                    if n > 205 else np.nan)

    out = {}
    for k, v in res.items():
        tol = 5e-3 if "RSI" in k else 1e-9 * max(1.0, float(np.nanmax(np.abs(c))))
        out[k] = (float(v), bool(v < tol) if not np.isnan(v) else None)
    return out


# =====================================================================
# 6. GHI KẾT QUẢ
# =====================================================================
def _write_excel(path: Path, summary, signals, quality, overview):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheets = [("Tong_hop", summary), ("Tin_hieu_chi_tiet", signals),
              ("Kiem_tra_du_lieu", quality), ("Thong_ke", overview)]
    fills = {"TÍCH CỰC": "C6EFCE", "TIÊU CỰC": "FFC7CE", "TRUNG LẬP": "FFEB9C",
             "Tích cực": "C6EFCE", "Tiêu cực": "FFC7CE", "Trung lập": "EDEDED"}
    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="DD/MM/YYYY") as xw:
        for name, frame in sheets:
            frame.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(frame.columns, start=1):
                letter = get_column_letter(i)
                width = min(60, max(10, len(str(col)) + 2,
                                    int(frame[col].astype(str).str.len().quantile(0.9)) + 2 if len(frame) else 10))
                ws.column_dimensions[letter].width = width
                ws.cell(row=1, column=i).font = Font(bold=True)
                ws.cell(row=1, column=i).alignment = Alignment(wrap_text=True, vertical="top")
                if col in ("Nhận định tổng hợp", "Tín hiệu", "Tín hiệu SMA", "Tín hiệu RSI",
                           "Tín hiệu MACD", "Tín hiệu Bollinger"):
                    for r in range(2, len(frame) + 2):
                        v = str(ws.cell(row=r, column=i).value)
                        key = "TRUNG LẬP" if v.startswith("TRUNG LẬP") else v
                        if key in fills:
                            ws.cell(row=r, column=i).fill = PatternFill("solid", fgColor=fills[key])
                if col in ("Giá đóng cửa", "SMA20", "SMA50", "EMA12", "EMA26", "BB_Upper", "BB_Middle", "BB_Lower"):
                    for r in range(2, len(frame) + 2):
                        ws.cell(row=r, column=i).number_format = "#,##0"
                if col in ("RSI14", "MACD", "MACD_Signal", "MACD_Hist"):
                    for r in range(2, len(frame) + 2):
                        ws.cell(row=r, column=i).number_format = "#,##0.00"


def main():
    p = argparse.ArgumentParser(description="Phân tích kỹ thuật nhiều mã cổ phiếu từ file OHLCV dạng dài.")
    p.add_argument("input", help="File CSV/XLSX (cột Date, Symbol, Open, High, Low, Close, Volume).")
    p.add_argument("--output-dir", default="output_technical", help="Thư mục lưu kết quả.")
    p.add_argument("--symbols", nargs="*", help="Chỉ phân tích các mã này (mặc định: tất cả).")
    p.add_argument("--keep-last", type=int, default=250,
                   help="Số phiên gần nhất mỗi mã lưu vào file chỉ báo (0 = toàn bộ, file rất lớn).")
    p.add_argument("--min-rows", type=int, default=MIN_ROWS, help="Số phiên tối thiểu để kết luận tín hiệu.")
    p.add_argument("--verify", action="store_true", help="Đối chiếu công thức với bản tính tay (tối đa 3 mã).")
    p.add_argument("--no-excel", action="store_true", help="Không xuất file Excel.")
    args = p.parse_args()

    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    df = load_data(args.input)
    if args.symbols:
        wanted = {s.upper() for s in args.symbols}
        unknown = sorted(wanted - set(df["Symbol"]))
        if unknown:
            print(f"[CẢNH BÁO] Không có trong dữ liệu: {', '.join(unknown)}")
        df = df[df["Symbol"].isin(wanted)]
    if df.empty:
        raise SystemExit("Không có dữ liệu để phân tích.")

    print(f"Đã đọc {len(df):,} dòng, {df['Symbol'].nunique():,} mã, "
          f"{df['Date'].min():%d/%m/%Y} -> {df['Date'].max():%d/%m/%Y}. "
          f"Ngày trùng đã loại: {df.attrs.get('duplicates_removed', 0)}.")

    quality = check_data(df)
    ind = add_indicators(df)
    summary, signals, assessment = analyze_all(ind, quality, args.min_rows)
    overview = market_overview(summary, quality)

    # --- ghi file ---
    summary.to_csv(out_dir / "technical_summary.csv", index=False, encoding="utf-8-sig")
    signals.to_csv(out_dir / "technical_signals_long.csv", index=False, encoding="utf-8-sig")
    quality.to_csv(out_dir / "data_quality_by_symbol.csv", index=False, encoding="utf-8-sig")
    (out_dir / "latest_assessment.txt").write_text(assessment, encoding="utf-8")
    keep = ind if args.keep_last <= 0 else ind.groupby("Symbol", sort=False).tail(args.keep_last)
    tag = "all" if args.keep_last <= 0 else f"last{args.keep_last}"
    keep.to_csv(out_dir / f"technical_indicators_{tag}.csv", index=False, encoding="utf-8-sig")
    if not args.no_excel:
        try:
            _write_excel(out_dir / "technical_report.xlsx", summary, signals, quality, overview)
        except ImportError:
            print("[CẢNH BÁO] Chưa cài openpyxl nên bỏ qua file Excel (pip install openpyxl).")

    # --- in kết quả ---
    print("\n=== THỐNG KÊ CHUNG ===")
    print(overview.to_string(index=False))
    bad = quality[quality["Chất lượng dữ liệu"] != "Tốt"]
    print(f"\n=== KIỂM TRA DỮ LIỆU: {len(bad)}/{len(quality)} mã có cảnh báo ===")
    for _, r in bad.head(15).iterrows():
        print(f" - {r['Mã']}: {r['Cảnh báo dữ liệu']}")
    if len(bad) > 15:
        print(f"   ... và {len(bad) - 15} mã khác (xem data_quality_by_symbol.csv)")
    show = ["Mã", "Ngày dữ liệu cuối", "Giá đóng cửa", "RSI14", "Nhận định tổng hợp", "Xu hướng", "Động lượng"]
    print("\n=== TÍN HIỆU (15 mã đầu) ===")
    view = summary[show].head(15).copy()
    view["Ngày dữ liệu cuối"] = view["Ngày dữ liệu cuối"].dt.strftime("%d/%m/%Y")
    view["RSI14"] = view["RSI14"].round(1)
    view["Giá đóng cửa"] = view["Giá đóng cửa"].map(lambda v: f"{v:,.0f}")
    print(view.to_string(index=False))

    if args.verify:
        print("\n=== ĐỐI CHIẾU CÔNG THỨC (sai số tuyệt đối lớn nhất) ===")
        for sym in list(ind["Symbol"].unique())[:3]:
            d = ind[ind["Symbol"] == sym].reset_index(drop=True)
            print(f"[{sym}] {len(d)} phiên")
            for k, (err, ok) in verify_formulas(d).items():
                flag = "bỏ qua (thiếu dữ liệu)" if ok is None else ("OK" if ok else "LỆCH")
                print(f"   {k:<24} {err:.2e}  {flag}")

    print(f"\nĐã lưu kết quả tại: {out_dir.resolve()}")
    print("Lưu ý: chỉ báo kỹ thuật có độ trễ, không phải khuyến nghị đầu tư.")


if __name__ == "__main__":
    main()
